"""Package H: dump the ccRCC mapping and clinical workbooks' sheets and first rows."""

import sys

import openpyxl

d = "/Users/rytis/ctx_kidney/data/raw/cptac_ccrcc/"
n = int(sys.argv[1]) if len(sys.argv) > 1 else 6
for f in [
    "S044_TMT10_Label_to_Sample_Mapping_File_CCRCC_r2_Jan2019.xlsx",
    "S050_S044_CPTAC_ccRCC_Discovery_Cohort_Clinical_Data_r4_Sept2019.xlsx",
]:
    b = openpyxl.load_workbook(d + f, read_only=True, data_only=True)
    for s in b.sheetnames:
        rows = list(b[s].iter_rows(values_only=True))
        print("==", f, s, len(rows), "rows")
        for r in rows[:n]:
            print(r[:40])
