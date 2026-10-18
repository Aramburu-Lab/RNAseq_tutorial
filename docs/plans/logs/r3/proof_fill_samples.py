"""R3 proof, the student's edit step: fill the starter sample table from PRIDE's sampletable.xlsx.

MaxQuant names a channel ``<n> plex<NN>``; the study names it ``TMT<n>.plex<NN>``. The plex is the
batch, the internal standard each plex's reference, and the RNA sample the TMT ID, the name the
paper's Table S2 RNA uses. No gather root holds that RNA, so ``ctx proteo`` cannot pair it yet.
Run through env.sh: python docs/plans/logs/r3/proof_fill_samples.py <starter csv> <sampletable.xlsx>
"""

import csv
import sys
from pathlib import Path

import openpyxl

starter, workbook = Path(sys.argv[1]), Path(sys.argv[2])
book = openpyxl.load_workbook(workbook, read_only=True)
rows = list(book[book.sheetnames[0]].iter_rows(values_only=True))
head = list(rows[0])
kind = {r[head.index("TMT_ID")]: r[head.index("sample.type")] for r in rows[1:]}
status = {r[head.index("TMT_ID")]: r[head.index("Status")] or "" for r in rows[1:]}
with starter.open() as handle:
    found = list(csv.reader(handle))
out = [["protein_sample", "rna_sample", "batch", "reference", "status", "sample_type"]]
for name, _, _ in found[1:]:
    channel, plex = name.split(" ")
    tmt = f"TMT{channel}.{plex}"
    sample_type = kind.get(tmt, "not listed")
    reference = "yes" if sample_type == "internal.standard" else ""
    rna = tmt if sample_type == "sample" else ""
    out.append([name, rna, plex, reference, status.get(tmt, ""), sample_type])
with starter.open("w", newline="") as handle:
    csv.writer(handle).writerows(out)
types = sorted({r[5] for r in out[1:]})
print(f"{starter.name}: {len(out) - 1} channels; sample types {types}")
print({t: sum(r[5] == t for r in out[1:]) for t in types})
