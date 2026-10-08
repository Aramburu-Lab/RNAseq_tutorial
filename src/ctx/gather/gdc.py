"""The one reader of a GDC project folder: counts, file metadata, clinical tables, annotations.

``ctx gather --gdc-dir`` is the only command that opens such a folder, as gather alone reads a
fetchngs outdir (D-006); every later module reads ``outputs/01_gather/`` and takes files only
through ``chosen()``. Fields are read by GDC's own column and key names, never by position or
by barcode shape, so a TCGA and a CPTAC folder pass alike; a TCGA aliquot barcode, where it
parses, adds vial, portion, plate and tissue source site. Every check runs before gather writes.
"""

import csv
import hashlib
import io
import json
import re
from collections.abc import Iterator, Mapping
from dataclasses import dataclass
from fractions import Fraction
from pathlib import Path
from typing import Any

import numpy as np
import polars as pl

from ctx.dataset import GdcSpec
from ctx.qc.registry import Metric

# The fetch's own file names, as the lab's GDC fetch script writes them.
METADATA = "gdc_files_star_counts.tsv"
MANIFEST = "gdc_manifest_star_counts.txt"
STATUS = "gdc_status.json"
RETRIEVAL = "_retrieval_log.tsv"
CASES = "gdc_cases_clinical.json"
ANNOTATIONS = "gdc_annotations.json"
QUERY = ".query.json"
# Our name for each GDC metadata column. ``sample`` is the aliquot: one counts file each.
FIELDS = {
    "file_id": "file_id", "file_name": "file_name", "md5sum": "md5sum", "file_size": "file_size",
    "case_id": "cases.0.case_id", "case": "cases.0.submitter_id",
    "sample_id": "cases.0.samples.0.submitter_id",
    "sample_type": "cases.0.samples.0.sample_type",
    "preservation": "cases.0.samples.0.preservation_method",
    "sample": "cases.0.samples.0.portions.0.analytes.0.aliquots.0.submitter_id",
}  # fmt: skip
PRIMARY_SITE = "cases.0.primary_site"
# A column for a second case, sample, portion, analyte or aliquot: a file of two would be
# counted for one alone.
SECOND = re.compile(r"(?:^|\.)(?:cases|samples|portions|analytes|aliquots)\.[1-9]")
N_ROWS = ("N_unmapped", "N_multimapping", "N_noFeature", "N_ambiguous")
COUNT_COLUMNS = ("unstranded", "stranded_first", "stranded_second")
GENE_COLUMNS = ("gene_id", "gene_name", "gene_type")
PAR_Y = "_PAR_Y"
# The call ``Metric.strandedness`` makes for each protocol's library.
CALLS = {"unstranded": "unstranded", "stranded_first": "forward", "stranded_second": "reverse"}
STRAND_METRIC = "gdc_strand_share"
FFPE = "FFPE"
# Why a file is not the one taken for its case and sample type, in the order the rule applies.
DROPS = ("ffpe", "protocol", "redacted", "not_first")
# Annotation entities that are biospecimens; the rest annotate other data types' files.
BIOSPECIMEN = ("case", "sample", "portion", "analyte", "aliquot")
REDACTION = "Redaction"
# A TCGA aliquot barcode: case (project-site-participant), sample type and vial, portion and
# analyte, plate, centre.
BARCODE = re.compile(
    r"(?P<case>[A-Z0-9]+-(?P<site>[A-Z0-9]{2})-[A-Z0-9]{4})-(?P<code>\d{2})(?P<vial>[A-Z])"
    r"-(?P<portion>\d{2})(?P<analyte>[A-Z])-(?P<plate>[A-Z0-9]{4})-(?P<centre>[A-Z0-9]{2})"
)
INT32_MAX = 2**31 - 1
CHUNK = 1 << 20


class GdcError(ValueError):
    """A GDC folder that cannot be used as it stands; the message lists every offender."""


@dataclass(frozen=True)
class Folder:
    """A GDC project folder's file list, checked against its manifest, release and log.

    Attributes:
        path: The folder.
        files: One row per counts file, by our names (``FIELDS``), in aliquot order.
        release: The data release ``gdc_status.json`` names.
        dropped_site: Files left out for another primary site.
    """

    path: Path
    files: pl.DataFrame
    release: str
    dropped_site: int = 0


@dataclass(frozen=True)
class Counts:
    """Every file's counts of the protocol's column, genes by files, and what each file held.

    Attributes:
        genes: gene_id (unversioned), gene_id_version, gene_name, gene_type, one row per gene.
        matrix: int32, genes by files in ``Folder.files`` order.
        stats: per file: sample, library_size, stranded_first, stranded_second and the
            protocol column's four ``N_`` rows.
        par_y: The ``_PAR_Y`` rows dropped, zero in every column of every file.
        rows: Rows in each file, ``N_`` rows included.
    """

    genes: pl.DataFrame
    matrix: np.ndarray
    stats: pl.DataFrame
    par_y: int
    rows: int


def _text(path: Path) -> str:
    """A file's text; a missing one is named."""
    if not path.is_file():
        raise GdcError(f"{path} is missing: not a GDC folder as the fetch writes it")
    return path.read_text(encoding="utf-8")


def _table(path: Path) -> list[dict[str, str]]:
    """A tab-separated file as rows by header name; CRLF ends a line, never a value."""
    rows = list(csv.DictReader(io.StringIO(_text(path), newline=""), delimiter="\t"))
    if bad := [
        v for r in rows for v in r.values() if isinstance(v, str) and ("\r" in v or "\n" in v)
    ]:
        raise GdcError(f"{path.name}: a value holds a line break: {bad[0]!r}")
    return rows


def _filters(node: Any) -> Iterator[tuple[str, list[str]]]:
    """Every (field, values) of a GDC query's filters."""
    if isinstance(node, Mapping):
        content = node.get("content")
        if isinstance(content, Mapping) and "field" in content:
            yield str(content["field"]), [str(v) for v in content.get("value", [])]
        for value in node.values():
            yield from _filters(value)
    elif isinstance(node, list):
        for value in node:
            yield from _filters(value)


def _queried(folder: Path, name: str, field: str) -> set[str]:
    """The values a fetch query filtered ``field`` on, from the query file beside ``name``."""
    query = json.loads(_text(folder / f"{name}{QUERY}"))
    return {v for f, values in _filters(query) if f.endswith(field) for v in values}


def _metadata(folder: Path, spec: GdcSpec) -> tuple[pl.DataFrame, int]:
    """The counts files' metadata by our names, kept to the declared primary site."""
    rows = _table(folder / METADATA)
    header = list(rows[0]) if rows else []
    problems = [f"{METADATA}: no column {c}" for c in FIELDS.values() if c not in header]
    problems += [f"{METADATA}: column {c} is for a second entity" for c in header
                 if SECOND.search(c)]  # fmt: skip
    if spec.primary_site and PRIMARY_SITE not in header:
        problems.append(f"{METADATA}: no {PRIMARY_SITE}, which a primary_site filter needs")
    if problems:
        raise GdcError("\n".join(problems))
    kept = [r for r in rows if not spec.primary_site or r[PRIMARY_SITE] == spec.primary_site]
    frame = pl.DataFrame({ours: [r[gdc] for r in kept] for ours, gdc in FIELDS.items()},
                         schema=dict.fromkeys(FIELDS, pl.String))  # fmt: skip
    frame = frame.with_columns(pl.col("file_size").cast(pl.Int64)).sort("sample")
    for column in ("file_id", "sample"):
        if frame[column].n_unique() != frame.height:
            twice = frame.filter(pl.col(column).is_duplicated())[column].unique().sort()
            raise GdcError(f"{METADATA}: {column} listed twice: {twice.to_list()[:5]}")
    return frame, len(rows) - len(kept)


def _manifest(folder: Path, files: pl.DataFrame) -> list[str]:
    """The manifest lists the same files with the same md5 and size."""
    listed = {r["id"]: r for r in _table(folder / MANIFEST)}
    problems = []
    for file_id, md5, size in files.select("file_id", "md5sum", "file_size").rows():
        row = listed.get(file_id)
        if row is None:
            problems.append(f"{MANIFEST}: no row for {file_id}")
        elif row["md5"] != md5 or int(row["size"]) != size:
            problems.append(f"{MANIFEST}: {file_id} differs from {METADATA} in md5 or size")
    return problems


def _release(folder: Path, spec: GdcSpec, files: pl.DataFrame) -> list[str]:
    """The status, every retrieval row and the file query name the declared release and query."""
    found = json.loads(_text(folder / STATUS)).get("data_release")
    problems = [] if found == spec.release else [f"{STATUS}: {found!r}, not {spec.release!r}"]
    log = _table(folder / RETRIEVAL)
    problems += [f"{RETRIEVAL}: {r['file']} retrieved from {r['gdc_release']!r}"
                 for r in log if r["gdc_release"] != spec.release][:5]  # fmt: skip
    logged = {r["file"] for r in log}
    needed = [METADATA, MANIFEST, CASES, ANNOTATIONS, STATUS]
    needed += [f"{i}/{n}" for i, n in files.select("file_id", "file_name").rows()]
    problems += [f"{RETRIEVAL}: no row for {n}" for n in needed if n not in logged][:5]
    for field, wanted in (("project_id", spec.project), ("workflow_type", spec.workflow)):
        if _queried(folder, METADATA, field) != {wanted}:
            problems.append(f"{METADATA}{QUERY}: {field} is not {wanted!r} alone")
    if _queried(folder, CASES, "project_id") != {spec.project}:
        problems.append(f"{CASES}{QUERY}: project_id is not {spec.project!r} alone")
    return problems


def read_folder(path: Path, spec: GdcSpec) -> Folder:
    """Read the file metadata and check it against the manifest, release, log and queries.

    Raises:
        GdcError: Listing every problem found.
    """
    files, dropped = _metadata(path, spec)
    if files.is_empty():
        raise GdcError(f"{path / METADATA}: no counts file of this project and site")
    problems = _manifest(path, files) + _release(path, spec, files)
    if problems:
        raise GdcError(f"{path}:\n  " + "\n  ".join(problems))
    return Folder(path, files, spec.release, dropped)


def _md5(path: Path) -> str:
    """MD5 of a file, read in chunks."""
    digest = hashlib.md5()  # noqa: S324 - GDC publishes MD5, the checksum compared here
    with path.open("rb") as handle:
        while block := handle.read(CHUNK):
            digest.update(block)
    return digest.hexdigest()


def verify(folder: Folder) -> None:
    """Every counts file is on disk with the MD5 GDC lists.

    Raises:
        GdcError: Listing every missing or mismatched file.
    """
    problems = []
    for file_id, name, md5 in folder.files.select("file_id", "file_name", "md5sum").rows():
        path = folder.path / file_id / name
        if not path.is_file():
            problems.append(f"missing {file_id}/{name}")
        elif _md5(path) != md5:
            problems.append(f"md5 mismatch {file_id}/{name}")
    if problems:
        raise GdcError(f"{len(problems)} of {folder.files.height} files fail:\n  "
                       + "\n  ".join(problems[:20]))  # fmt: skip


def _one_file(path: Path, spec: GdcSpec) -> pl.DataFrame:
    """One counts file's rows, after its gene-model line is checked."""
    with path.open(encoding="utf-8") as handle:
        first = handle.readline().rstrip("\r\n")
    if first != f"# gene-model: {spec.gene_model}":
        raise GdcError(f"{path.name}: first line {first!r}, not '# gene-model: {spec.gene_model}'")
    schema = {**dict.fromkeys(GENE_COLUMNS, pl.String), **dict.fromkeys(COUNT_COLUMNS, pl.Int64)}
    try:
        return pl.read_csv(path, separator="\t", comment_prefix="#", columns=list(schema),
                           schema_overrides=schema)  # fmt: skip
    except (pl.exceptions.PolarsError, ValueError) as error:
        raise GdcError(f"{path.name}: {error}") from error


def _genes(reference: pl.DataFrame) -> pl.DataFrame:
    """The genes of the reference file: _PAR_Y rows out, then versions stripped, unique."""
    kept = reference.filter(~pl.col("gene_id").str.ends_with(PAR_Y))
    genes = kept.select(
        gene_id=pl.col("gene_id").str.replace(r"\.\d+$", ""),
        gene_id_version="gene_id",
        gene_name="gene_name",
        gene_type="gene_type",
    )
    if genes["gene_id"].n_unique() != genes.height:
        twice = genes.filter(pl.col("gene_id").is_duplicated())["gene_id"].unique().sort()
        raise GdcError(
            f"unversioned gene IDs not unique once _PAR_Y rows leave: {twice[:5].to_list()}"
        )
    return genes


def read_counts(folder: Folder, spec: GdcSpec) -> Counts:
    """Every counts file, checked against the first: gene model, gene order, N_ rows, _PAR_Y.

    _PAR_Y rows must be zero in every count column of every file before they leave, since
    their unversioned IDs repeat their X copies'. The matrix holds the protocol's column.

    Raises:
        GdcError: On the first file that fails a check.
    """
    reference: pl.DataFrame | None = None
    matrix = np.empty(0, dtype=np.int32)
    stats = []
    gene_rows = np.empty(0, dtype=bool)
    for j, (sample, file_id, name) in enumerate(
        folder.files.select("sample", "file_id", "file_name").rows()
    ):
        frame = _one_file(folder.path / file_id / name, spec)
        is_n = frame["gene_id"].str.starts_with("N_")
        if sorted(frame.filter(is_n)["gene_id"]) != sorted(N_ROWS):
            raise GdcError(f"{name}: N_ rows {frame.filter(is_n)['gene_id'].to_list()}")
        if reference is None:
            reference = frame.filter(~is_n)
            gene_rows = ~reference["gene_id"].str.ends_with(PAR_Y).to_numpy()
            matrix = np.zeros((int(gene_rows.sum()), folder.files.height), dtype=np.int32)
        body = frame.filter(~is_n)
        if not body["gene_id"].equals(reference["gene_id"]):
            raise GdcError(f"{name}: gene order differs from {folder.files['file_name'][0]}")
        par_y = body.filter(pl.col("gene_id").str.ends_with(PAR_Y)).select(COUNT_COLUMNS)
        if par_y.height and int(par_y.select(pl.sum_horizontal(pl.all()).sum()).item()):
            raise GdcError(f"{name}: a {PAR_Y} row holds counts; its unversioned ID would repeat")
        counts = body[spec.protocol].to_numpy()[gene_rows]
        if counts.min() < 0 or counts.max() > INT32_MAX:
            raise GdcError(f"{name}: a count outside int32")
        matrix[:, j] = counts
        n_rows = dict(frame.filter(is_n).select("gene_id", spec.protocol).rows())
        sums = body.filter(~pl.col("gene_id").str.ends_with(PAR_Y)).select(
            pl.col("stranded_first", "stranded_second").sum()
        )
        stats.append({"sample": sample, "library_size": int(counts.sum(dtype=np.int64)),
                      "stranded_first": int(sums["stranded_first"][0]),
                      "stranded_second": int(sums["stranded_second"][0]),
                      **{n.lower(): int(n_rows[n]) for n in N_ROWS}})  # fmt: skip
    assert reference is not None  # noqa: S101 - read_folder refuses a folder without files
    return Counts(_genes(reference), matrix, pl.DataFrame(stats),
                  int((~gene_rows).sum()), reference.height + len(N_ROWS))  # fmt: skip


def strand(stats: pl.DataFrame, metric: Metric) -> pl.DataFrame:
    """Each file's first-strand share over genes and the library type it shows.

    The call is the QC registry's ``strandedness`` on the two sums as a Fraction, exact at a
    band's edge (lessons.md, 2026-10-01); a file with neither sum is undetermined.
    """
    calls = [metric.strandedness(Fraction(f, f + s)) if f + s else "undetermined"
             for f, s in stats.select("stranded_first", "stranded_second").rows()]  # fmt: skip
    share = pl.col("stranded_first") / (pl.col("stranded_first") + pl.col("stranded_second"))
    return stats.with_columns(first_strand_share=share, strand_call=pl.Series(calls))


def barcode(sample: str, case: str) -> dict[str, str]:
    """Vial, portion, plate and tissue source site of a TCGA aliquot barcode of this case.

    Empty for any other identifier: a CPTAC aliquot carries none of them.
    """
    found = BARCODE.fullmatch(sample)
    if found is None or found["case"] != case:
        return dict.fromkeys(("vial", "portion", "plate", "site"), "")
    return {k: found[k] for k in ("vial", "portion", "plate", "site")}


def annotations(path: Path, files: pl.DataFrame) -> pl.DataFrame:
    """GDC's annotations of this folder's cases' biospecimens, and the counts files each hits.

    A case, sample or aliquot annotation hits the files of that entity; a portion or analyte
    one, the aliquots under it, whose TCGA barcodes extend its own. Only a redaction that hits
    a file removes it (CONVENTIONS.md, Data and scope); every other annotation is logged.

    Returns:
        case, classification, category, entity_type, entity, files hit and the action.
    """
    hits = json.loads(_text(path))["data"]["hits"]
    cases = set(files["case"])
    keys = {"case": "case", "sample": "sample_id", "aliquot": "sample"}
    by_key: dict[str, dict[str, list[str]]] = {k: {} for k in keys}
    for row in files.select(*dict.fromkeys(keys.values())).iter_rows(named=True):
        for kind, column in keys.items():
            by_key[kind].setdefault(row[column], []).append(row["sample"])
    rows = []
    for hit in hits:
        case, kind = hit.get("case_submitter_id"), hit.get("entity_type")
        if case not in cases or kind not in BIOSPECIMEN:
            continue
        entity = str(hit.get("entity_submitter_id"))
        if kind in by_key:
            touched = list(by_key[kind].get(entity, []))
        else:
            touched = [s for s in by_key["case"][case] if s.startswith(f"{entity}-")]
        redacts = hit.get("classification") == REDACTION and bool(touched)
        rows.append((case, str(hit.get("classification")), str(hit.get("category")), kind,
                     entity, len(touched), ";".join(sorted(touched)),
                     "drops its files" if redacts else "logged"))  # fmt: skip
    schema = {**dict.fromkeys(("case", "classification", "category", "entity_type", "entity"),
                              pl.String), "files": pl.Int64, "samples": pl.String,
              "action": pl.String}  # fmt: skip
    return pl.DataFrame(rows, schema=schema, orient="row").sort("case", "category", "entity")


def choose(sheet: pl.DataFrame, protocol: str) -> pl.DataFrame:
    """One file per case and sample type, by the approved rule (decision 1, 2026-10-08).

    In this order: drop FFPE, drop libraries of another protocol, drop redacted files, then
    take the first aliquot ID of what is left. Every file gets ``chosen`` and the reason it is
    not, ``drop_reason``; the tie-break runs last, so an FFPE or stranded aliquot that sorts
    first cannot take a case.
    """
    reason = (
        pl.when(pl.col("preservation") == FFPE).then(pl.lit("ffpe"))
        .when(pl.col("strand_call") != CALLS[protocol]).then(pl.lit("protocol"))
        .when(pl.col("redacted")).then(pl.lit("redacted"))
        .otherwise(pl.lit(""))
    )  # fmt: skip
    marked = sheet.with_columns(drop_reason=reason)
    rank = pl.col("sample").rank("ordinal").over("case", "sample_type", "drop_reason")
    first = (pl.col("drop_reason") == "") & (rank == 1)
    return marked.with_columns(
        chosen=first,
        drop_reason=pl.when(first | (pl.col("drop_reason") != ""))
        .then("drop_reason")
        .otherwise(pl.lit("not_first")),
    )


def chosen(sheet: pl.DataFrame, sample_type: str | None = None) -> pl.DataFrame:
    """The one door: the files taken, one per case and sample type, optionally of one type.

    Every consumer of gather's GDC sheet selects through this function; a sheet that holds only
    the chosen rows passes through unchanged.
    """
    kept = sheet.filter(pl.col("chosen"))
    return kept if sample_type is None else kept.filter(pl.col("sample_type") == sample_type)


def _primary(diagnoses: list[dict[str, Any]]) -> dict[str, Any]:
    """The diagnosis GDC marks as the primary disease, else the first."""
    marked = [d for d in diagnoses if d.get("diagnosis_is_primary_disease")]
    return (marked or diagnoses or [{}])[0]


def clinical(path: Path) -> tuple[pl.DataFrame, pl.DataFrame, pl.DataFrame]:
    """GDC's cases: one row per case, then treatments and molecular tests, one row each.

    Sex is ``demographic.sex_at_birth``: Data Release 46.0 has no ``gender`` (plan, 1.1).
    """
    hits = json.loads(_text(path))["data"]["hits"]
    cases, treatments, tests = [], [], []
    for hit in hits:
        case, demographic = hit["submitter_id"], hit.get("demographic") or {}
        primary = _primary(hit.get("diagnoses") or [])
        cases.append({
            "case": case, "case_id": hit.get("case_id", ""),
            "primary_site": hit.get("primary_site") or "",
            "disease_type": hit.get("disease_type") or "",
            "sex_at_birth": demographic.get("sex_at_birth") or "",
            "age_at_index": demographic.get("age_at_index"),
            "vital_status": demographic.get("vital_status") or "",
            "primary_diagnosis": primary.get("primary_diagnosis") or "",
            "ajcc_pathologic_stage": primary.get("ajcc_pathologic_stage") or "",
        })  # fmt: skip
        for diagnosis in hit.get("diagnoses") or []:
            treatments += [(case, t.get("treatment_type") or "",
                            t.get("treatment_or_therapy") or "")
                           for t in diagnosis.get("treatments") or []]  # fmt: skip
        for follow_up in hit.get("follow_ups") or []:
            tests += [
                (
                    case,
                    m.get("gene_symbol") or "",
                    m.get("molecular_analysis_method") or "",
                    m.get("test_result") or "",
                )
                for m in follow_up.get("molecular_tests") or []
            ]
    schema = {"age_at_index": pl.Int64}
    return (pl.DataFrame(cases, schema_overrides=schema).sort("case"),
            pl.DataFrame(treatments, schema=["case", "treatment_type", "treatment_or_therapy"],
                         orient="row").sort("case", "treatment_type"),
            pl.DataFrame(tests, schema=["case", "gene_symbol", "method", "result"],
                         orient="row").sort("case", "gene_symbol", "method", "result"))  # fmt: skip


def _cell(value: Any) -> str:
    """A workbook cell as the text it shows: whole numbers without a decimal point."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def cdr(path: Path, cdr_type: str) -> pl.DataFrame:
    """The TCGA-CDR rows of one ``type``, every column as the workbook's text, verbatim.

    Parsing (days, missing tokens, stage) is ``ctx.clinical.harmonise``'s; gather copies.
    The workbook's unnamed first column, a row number, is left out.

    Raises:
        GdcError: If openpyxl is not installed, or the sheet or its columns are missing.
    """
    try:
        import openpyxl  # noqa: PLC0415 - only a CDR read needs it
    except ImportError as error:
        raise GdcError("reading the TCGA-CDR needs openpyxl ==3.1.5, pinned in pixi.toml; "
                       "run pixi install") from error  # fmt: skip
    if not path.is_file():
        raise GdcError(f"{path} is missing; [gdc] cdr_file names it")
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if "TCGA-CDR" not in book.sheetnames:
            raise GdcError(f"{path.name}: no TCGA-CDR sheet")
        rows = book["TCGA-CDR"].iter_rows(values_only=True)
        header = [_cell(h) for h in next(rows)]
        named = [i for i, h in enumerate(header) if h]
        if "type" not in header or "bcr_patient_barcode" not in header:
            raise GdcError(f"{path.name}: TCGA-CDR has no type or bcr_patient_barcode column")
        kind = header.index("type")
        kept = [[_cell(r[i]) for i in named] for r in rows if _cell(r[kind]) == cdr_type]
    finally:
        book.close()
    frame = pl.DataFrame(kept, schema=dict.fromkeys((header[i] for i in named), pl.String),
                         orient="row")  # fmt: skip
    if frame.is_empty():
        raise GdcError(f"{path.name}: no TCGA-CDR row of type {cdr_type!r}")
    return frame
