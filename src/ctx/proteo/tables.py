"""The protein readers: a GCT member of a data-freeze archive, and a copy-number workbook.

Both end in a ``Layer``: one row per gene symbol, a missing value as NaN, never as zero
(CONVENTIONS.md, proteomics). ``one_per_gene`` is the single place where protein groups become
genes, and it reports what it dropped: rows with no symbol, groups naming several genes (shared
peptides), and every extra row of a gene measured twice (isoforms).
"""

import re
import tarfile
from dataclasses import dataclass
from pathlib import Path

import numpy as np
import polars as pl

from ctx.proteo import ProteoError

# GCT 1.3: version, dimensions, header, column annotations, data (clue.io connectopedia).
GCT_VERSION = "#1.3"
GCT_MISSING = ("NA", "NaN", "")
# The freeze's own gene symbol column among the row descriptors.
GCT_SYMBOL = "geneSymbol"
# Rieckmann's Table 6: the protein group and its genes, as MaxQuant names them.
GROUP, GENES = "Majority protein IDs", "Gene names"
# MaxQuant joins several genes of one group with a semicolon.
SEVERAL = ";"


@dataclass(frozen=True)
class Collapse:
    """What ``one_per_gene`` did: every row read is kept or counted under one reason.

    Attributes:
        rows: Rows read.
        no_symbol: Rows with no gene symbol.
        several_genes: Rows whose group names more than one gene.
        extra_rows: Further rows of a gene already kept (isoforms or split groups).
        genes_twice: Genes measured by more than one row.
        kept: Genes kept, one row each.
    """

    rows: int
    no_symbol: int
    several_genes: int
    extra_rows: int
    genes_twice: int
    kept: int


@dataclass(frozen=True)
class Layer:
    """One measurement layer: genes by samples, NaN where a value is missing.

    Attributes:
        genes: Gene symbols, one per row, unique.
        samples: Sample names, one per column.
        values: float64, genes by samples.
        collapse: How the rows became genes; None for a layer read one row per gene.
        zeros: Cells written as 0 in the source and read as missing.
    """

    genes: tuple[str, ...]
    samples: tuple[str, ...]
    values: np.ndarray
    collapse: Collapse | None = None
    zeros: int = 0

    def row(self) -> dict[str, int]:
        """Gene symbol to row index."""
        return {g: i for i, g in enumerate(self.genes)}


def one_per_gene(symbols: list[str | None], values: np.ndarray) -> tuple[list[int], Collapse]:
    """The row kept for each gene, in file order, and what was dropped.

    A row naming no gene or several is dropped; of a gene's rows, the one quantified in the
    most samples is kept, the earlier in the file on a tie.

    Args:
        symbols: Each row's gene symbol, several joined by ``SEVERAL``; None or empty for none.
        values: The rows' values, NaN where missing.

    Returns:
        The kept row indices in file order, and the report.
    """
    present = (~np.isnan(values)).sum(axis=1)
    best: dict[str, int] = {}
    seen: dict[str, int] = {}
    no_symbol = several = 0
    for i, symbol in enumerate(symbols):
        name = (symbol or "").strip()
        if not name:
            no_symbol += 1
            continue
        if SEVERAL in name:
            several += 1
            continue
        seen[name] = seen.get(name, 0) + 1
        if name not in best or present[i] > present[best[name]]:
            best[name] = i
    kept = sorted(best.values())
    extra = sum(n - 1 for n in seen.values())
    report = Collapse(len(symbols), no_symbol, several, extra,
                      sum(n > 1 for n in seen.values()), len(kept))  # fmt: skip
    return kept, report


def _member(archive: Path, name: str) -> bytes:
    """One file of a tar archive, found by its base name.

    Raises:
        ProteoError: If the archive is missing or holds no such file, or holds it twice.
    """
    if not archive.is_file():
        raise ProteoError(f"{archive} is missing; run ctx setup with the dataset file")
    with tarfile.open(archive) as tar:
        found = [m for m in tar.getmembers() if m.isfile() and Path(m.name).name == name]
        if len(found) != 1:
            raise ProteoError(f"{archive.name}: {len(found)} members named {name}")
        handle = tar.extractfile(found[0])
        assert handle is not None  # noqa: S101 - a regular member always has content
        return handle.read()


def read_gct(archive: Path, member: str) -> tuple[Layer, pl.DataFrame]:
    """A GCT 1.3 table from the archive: one row per gene, and its column annotations.

    Raises:
        ProteoError: If the member is not GCT 1.3, its dimensions disagree with its body, or it
            has no ``geneSymbol`` row descriptor.
    """
    lines = _member(archive, member).decode("utf-8").splitlines()
    if not lines or lines[0].strip() != GCT_VERSION:
        raise ProteoError(f"{member}: not {GCT_VERSION}")
    n_rows, n_cols, n_desc, n_ann = (int(x) for x in lines[1].split("\t")[:4])
    header = lines[2].split("\t")
    samples = header[1 + n_desc :]
    body = lines[3 + n_ann :]
    if len(samples) != n_cols or len(body) != n_rows:
        raise ProteoError(f"{member}: {len(body)} rows and {len(samples)} samples, "
                          f"not the {n_rows} and {n_cols} its second line declares")  # fmt: skip
    if GCT_SYMBOL not in header[1 : 1 + n_desc]:
        raise ProteoError(f"{member}: no {GCT_SYMBOL} row descriptor")
    symbol_at = header.index(GCT_SYMBOL)
    cells = [line.split("\t") for line in body]
    values = np.array([[np.nan if c in GCT_MISSING else float(c) for c in r[1 + n_desc :]]
                       for r in cells], dtype=np.float64)  # fmt: skip
    kept, report = one_per_gene([r[symbol_at] for r in cells], values)
    annotation = pl.DataFrame(
        {row[0]: row[1 + n_desc :] for row in (line.split("\t") for line in lines[3 : 3 + n_ann])}
    ).with_columns(pl.Series("sample", samples))
    genes = tuple(cells[i][symbol_at].strip() for i in kept)
    return Layer(genes, tuple(samples), values[kept], report), annotation


def read_copies(path: Path, tested: str, reference: str) -> tuple[Layer, int, int]:
    """Copy numbers of the two groups' columns, log2, one row per gene; 0 and blank missing.

    MaxQuant writes 0 for a protein not quantified in a sample: it is read as missing.

    Args:
        path: The workbook.
        tested: Regular expression matching the tested group's column names.
        reference: The same for the reference group.

    Returns:
        The layer (tested columns first), and how many columns each group has.

    Raises:
        ProteoError: If openpyxl is missing, the workbook lacks the gene column, or a group
            matches no column.
    """
    try:
        import openpyxl  # noqa: PLC0415 - only the copy-number table needs it
    except ImportError as error:
        raise ProteoError("reading the copy numbers needs openpyxl ==3.1.5, pinned in pixi.toml"
                          ) from error  # fmt: skip
    if not path.is_file():
        raise ProteoError(f"{path} is missing; run ctx setup with the dataset file")
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        rows = book[book.sheetnames[0]].iter_rows(values_only=True)
        header = [str(h) for h in next(rows)]
        picks = [[i for i, h in enumerate(header) if re.search(p, h)] for p in (tested, reference)]
        if GENES not in header or not all(picks):
            raise ProteoError(f"{path.name}: needs {GENES!r} and columns for both groups")
        at = header.index(GENES)
        columns = picks[0] + picks[1]
        symbols, raw = [], []
        for r in rows:
            symbols.append(None if r[at] is None else str(r[at]))
            raw.append([r[i] for i in columns])
    finally:
        book.close()
    cells = np.array([[np.nan if v is None or v == "" else float(v) for v in row] for row in raw])
    if (cells < 0).any():
        raise ProteoError(f"{path.name}: a negative copy number")
    zeros = int((cells == 0).sum())
    with np.errstate(divide="ignore"):
        values = np.where(cells > 0, np.log2(np.where(cells > 0, cells, 1.0)), np.nan)
    kept, report = one_per_gene(symbols, values)
    layer = Layer(tuple(str(symbols[i]).strip() for i in kept),
                  tuple(header[i] for i in columns), values[kept], report, zeros)  # fmt: skip
    return layer, len(picks[0]), len(picks[1])
